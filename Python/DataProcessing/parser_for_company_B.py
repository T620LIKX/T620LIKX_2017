#!/usr/bin/env python3

import psycopg2
import openpyxl

def connect_to_database(host, dbname, username, pw):
    conn_string = "host='{}' dbname='{}' user='{}' password='{}'".format(host, dbname, username, pw)
    try:
        conn = psycopg2.connect(conn_string)
    except psycopg2.OperationalError as e:
        print('Connection failed!')
        print('Error message:', e)
        exit()
    cursor = conn.cursor()
    return cursor, conn


createtablestatement = """create table callcenter_b
(date_and_time timestamp,
 call_center_name text,
 calls_queued int,
 calls_escaped int,
 calls_abandoned int,
 calls_presented int,
 calls_answered int,
 calls_answered_in_180_secs int,
 perc_calls_answered_in_180_secs real,
 perc_answered_calls_answered_in_180_secs real,
 calls_overflowed_time int,
 calls_bounced int,
 calls_bounced_transferred int,
 calls_transferred int,
 calls_stranded int,
 calls_stranded_unavailable int);
 """

print('Connecting to the database.')
cursor, conn = connect_to_database('localhost', 'likanx', 'postgres', 'postgres')

print('Create the callcenter_b table.')
# Create the table, first check if the table is there.
cursor.execute("select exists(select * from information_schema.tables where table_name=%s)", ('callcenter_b',))
if cursor.fetchone()[0] == False:
    cursor.execute(createtablestatement)
    conn.commit()

excelfiles = ['Símtöl 2016 - 1 & 2.xlsx', 'Símtöl 2016 - 3 & 4.xlsx']

for excelfile in excelfiles:
    print('Opening the Excel files: {}'.format(excelfile))
    wb = openpyxl.load_workbook(filename=excelfile)

    alldata = []

    for sheetname in wb.sheetnames:
        print('Parsing worksheet: {}'.format(sheetname))
        ws = wb.get_sheet_by_name(sheetname)

        data = ws.values
        data = list(data)

        if len(data) > 10:
            i = 0
            while data[i][0] != 'Call Center Activity':
                i += 1

            i = i+2  # skip over 'call center activity' and the header line

            last_timedate = ''
            sheetdata = []
            while data[i][0] != 'Report Summary':
                data[i] = list(data[i])
                if data[i][0] != None:
                    last_timedate = data[i][0]
                else:
                    data[i][0] = last_timedate

                if not (data[i][1] == 'Summary' or data[i][2] == None):
                    for x in range(len(data[i])):
                        if type(data[i][x]) == str and '%' in data[i][x]:
                            data[i][x] = float(data[i][x].strip('%'))/100.0
                    sheetdata.append(data[i])

                i += 1

            alldata = alldata + sheetdata

    numberofrowstoinsert = 500
    counter = 0

    print('Insert the data into the database.')

    insertstring = """insert into callcenter_b (
     date_and_time,
     call_center_name,
     calls_queued,
     calls_escaped,
     calls_abandoned,
     calls_presented,
     calls_answered,
     calls_answered_in_180_secs,
     perc_calls_answered_in_180_secs,
     perc_answered_calls_answered_in_180_secs,
     calls_overflowed_time,
     calls_bounced,
     calls_bounced_transferred,
     calls_transferred,
     calls_stranded,
     calls_stranded_unavailable) values """
    values = []

    for i in alldata:
        values.append( i )
        counter += 1

        if counter == numberofrowstoinsert:
            args_str = b','.join(cursor.mogrify("(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)", x) for x in values)
            cursor.execute(insertstring + args_str.decode('utf-8'))
            values = []
            counter = 0

    if len(values) > 0:
        args_str = b','.join(cursor.mogrify("(%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)", x) for x in values)
        cursor.execute(insertstring + args_str.decode('utf-8'))
        values = []
        counter = 0

conn.commit()
cursor.close()
conn.close()

